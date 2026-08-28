"""Deterministic loader for tet-spec TCR tracking sheets -> TcrClonotype + observations.

Motivation (2026-07-04): the live LLM extractor is (1) NON-DETERMINISTIC on the TCR lanes — the same
paper yields wildly different clonotype counts run-to-run (Hu 33479501: 0 -> 39 -> 20; 39972124: 0 -> 12)
— and (2) budget-shortcuts dense tables, capturing a clone's origin-frequency snapshot but NOT its full
per-timepoint trajectory. Both failures are structural to LLM extraction of large, clean, columnar
supplements. This adapter is the alternative for that class of table: a PURE FUNCTION of the sheet, so
the output is byte-identical on every run, and it reads EVERY timepoint column, so the trajectory is
complete.

This is NOT an `add_table` mapping — that DSL is a flat column->field copy and cannot UNPIVOT a wide
sheet (one CDR3b row × N timepoint columns) into the LONG, NESTED `observations` list. Hence a bespoke
adapter, in the spirit of the schema's existing paper-specific table adapters.

Shape handled (Hu 2021 / 33479501, "Suppl 8 Pt<N> Tet-spec TCR<a|b> track" sheets):
    col 0            = the clone identity — EITHER a bare CDR3 ('AGQGFNQPQHF', Pt2) OR the paper's
                       comma-joined 'V,J,CDR3' token ('TRBV5-1,TRBJ2-1,CASSLDRNNEQFF', Pt3-Pt6)
    col 1            = "Tetramer-Specific TCR?" (1 = yes)   <- we keep ONLY these rows
    col 2..k         = frequency of that clone at each timepoint (fraction), header row names them
                       (Pre-Vax, Week 3/4/8/12/16/20/24, Post-pembro, Long-term — the set VARIES by
                       patient, and the label row is padded with blank trailing columns)
Both loci occur (TCR⍺ and TCRβ track sheets per patient); `locus` is inferred from the identity
column's header when not passed.

Wired into the live extractor (2026-07-26) as the MCP tool `add_tcr_track` ->
`agent_core.load_tcr_track_into`, which validates, merge-dedupes by (patient, chain identity) and LOCKS
the `tcr_clonotypes` lane. The paired single-cell one-row-per-CELL sheet is `tcr_reactive_adapter` /
`add_reactive_tcr` instead; both write the same lane.
"""
from __future__ import annotations

import pathlib
import re

from . import schema

# Timepoint label -> the reused Evidence TimepointPhase (a documented mapping choice, not inference from
# data). Unknown labels keep the verbatim label with timepoint_phase=None.
_TIMEPOINT_PHASE = {
    "pre-vax": "pre_vaccine", "pre vax": "pre_vaccine", "prevax": "pre_vaccine", "baseline": "pre_vaccine",
    "post-pembro": "post_vaccine", "post pembro": "post_vaccine",
    "long-term": "memory", "long term": "memory", "longterm": "memory",
}
# Any 'Week N' is on-treatment — the WEEKS DIFFER BY PATIENT in one supplement (Pt2 runs 4/12/16/24,
# Pt3-Pt6 run 3/8/12/16/20/24), so this is a rule, not an enumeration that silently drops the odd week.
_WEEK = re.compile(r"^week\s*\d+$")
_JUNCTION = re.compile(r"^C[ACDEFGHIKLMNPQRSTVWY]{2,38}[FW]$")   # IMGT junction: C…F/W
_AA = re.compile(r"^[ACDEFGHIKLMNPQRSTVWY]{4,40}$")              # the schema's CDR3 pattern
_V_GENE = re.compile(r"^TR[ABGD]V", re.I)
_J_GENE = re.compile(r"^TR[ABGD]J", re.I)


def _phase(label):
    l = label.strip().lower()
    return "on_treatment" if _WEEK.match(l) else _TIMEPOINT_PHASE.get(l)


def _infer_locus(header_cell):
    """The identity column's header names the chain ('CDR3β' / 'TCR⍺'). Returns 'TRB' / 'TRA' / None."""
    s = str(header_cell or "").strip().lower()
    if "β" in s or "beta" in s or re.search(r"(tcr|cdr3)\s*b\b", s):
        return "TRB"
    if "⍺" in s or "α" in s or "alpha" in s or re.search(r"(tcr|cdr3)\s*a\b", s):
        return "TRA"
    return None


def _split_identity(cell):
    """The identity cell -> (cdr3_aa, v_call, j_call). Two forms occur in ONE supplement: a bare CDR3
    (Pt2) and the paper's comma-joined 'V,J,CDR3' token (Pt3-Pt6, e.g. 'TRBV6-2/3/5/6,TRBJ1-5,CASS…').
    Gene calls are kept VERBATIM (the schema deliberately does not regex them to IMGT syntax); the CDR3
    is the last comma field, so an unfamiliar leading field never gets mistaken for the sequence."""
    parts = [p.strip() for p in str(cell).split(",") if p.strip()]
    if not parts:
        return None, None, None
    cdr3 = parts[-1].upper()
    if not _AA.match(cdr3):
        return None, None, None                    # not a CDR3 at all -> caller drops the row
    v = next((p for p in parts[:-1] if _V_GENE.match(p)), None)
    j = next((p for p in parts[:-1] if _J_GENE.match(p)), None)
    return cdr3, v, j


def _num(v):
    """A cell as a float, or None if it isn't a number (blank / label / '-')."""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().rstrip("%")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def parse_tetraspec_track(header_row, data_rows, *, patient_paper_id, locus=None,
                          tissue="blood", sorted_compartment="whole_pbmc",
                          cdr3_col=0, flag_col=1, first_timepoint_col=2):
    """PURE FUNCTION: (header, rows) -> list[schema.TcrClonotype]. Deterministic — same input, same
    output (stable row order, no randomness). One kept row (Tetramer-Specific=1) -> one single-chain
    clonotype whose `observations` are the non-blank per-timepoint frequencies, UNPIVOTED from the wide
    columns. observations_identity='paper_explicit' (the tracking table IS the paper asserting the same
    clone across timepoints). Cross-linking to a paired hero clone is left to a CDR3 query (decision 1b).

    `locus` ('TRB'/'TRA') is inferred from the identity column's header when not passed; a sheet whose
    header names neither chain raises rather than guessing (a wrong locus is silent bad data).
    """
    locus = locus or _infer_locus(header_row[cdr3_col] if cdr3_col < len(header_row) else None)
    if locus not in ("TRA", "TRB"):
        raise ValueError(
            f"cannot tell which chain this track sheet holds from its identity-column header "
            f"{header_row[cdr3_col]!r} — pass locus='TRB' or 'TRA' explicitly")
    pairing = "trb_only" if locus == "TRB" else "tra_only"
    timepoint_labels = [str(c).strip() if c is not None else "" for c in header_row[first_timepoint_col:]]
    out = []
    for r in data_rows:
        if len(r) <= flag_col:
            continue
        raw_identity = r[cdr3_col]
        if not isinstance(raw_identity, str) or not raw_identity.strip():
            continue
        raw_identity = raw_identity.strip()
        cdr3, v_call, j_call = _split_identity(raw_identity)
        if not cdr3:
            continue
        if _num(r[flag_col]) != 1.0:          # keep ONLY the tetramer-specific clones (the paper's flag)
            continue
        observations = []
        for j, label in enumerate(timepoint_labels):
            col = first_timepoint_col + j
            if col >= len(r):
                break
            freq = _num(r[col])
            if freq is None:
                continue                       # clone not detected at this timepoint -> no observation
            observations.append(schema.TcrObservation(
                tissue=tissue, sorted_compartment=sorted_compartment,
                timepoint_phase=_phase(label),
                timepoint_label=label or None,
                frequency_value=freq, frequency_basis="fraction",
                raw=f"{label}: {freq}",
            ))
        if not observations:
            continue
        # nomenclature: a C…F/W sequence is a junction (incl. flanks); else CDR3-proper.
        is_junction = bool(_JUNCTION.match(cdr3))
        chain = schema.TcrChain(
            locus=locus,
            junction_aa=cdr3 if is_junction else None,
            cdr3_aa=None if is_junction else cdr3,
            cdr3_scheme="imgt_junction_aa" if is_junction else "cdr3_aa_no_flank",
            v_call=v_call, j_call=j_call,
            raw=raw_identity if raw_identity.upper() != cdr3 else None,   # lossless V,J,CDR3 token
        )
        quoted = f"tet-spec track {raw_identity}: " + "; ".join(o.raw for o in observations)
        if len(quoted) > 2000:                 # ShortText cap — a very long track never fails the load
            quoted = quoted[:1997] + "..."
        out.append(schema.TcrClonotype(
            paper_local_id=f"{patient_paper_id}_track_{locus}_{cdr3}",
            patient_paper_id=patient_paper_id,
            receptor_modality="bulk_tcr_seq",
            chain_pairing=pairing,
            chains=[chain],
            observations_identity="paper_explicit",
            observations=observations,
            quoted_text=quoted,
            section_ref="Suppl Dataset 8 (tet-spec track)",
        ))
    return out


def load_track_sheet(xlsx_path, sheet, *, patient_paper_id, header_row_index=1, **kw):
    """I/O wrapper: stream one track sheet (read_only for the 257MB Hu supplement) and parse it.
    header_row_index is the 0-based row that carries the timepoint labels (row above the data)."""
    import openpyxl
    big = pathlib.Path(xlsx_path).stat().st_size > 50_000_000
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=big)
    try:
        ws = wb[sheet]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        if big:
            wb.close()
    header = rows[header_row_index]
    data = rows[header_row_index + 1:]
    return parse_tetraspec_track(header, data, patient_paper_id=patient_paper_id, **kw)
