"""PROTOTYPE — deterministic loader for class-I/II prediction / epitope-manifest supplementary tables.

Motivation (stability lab, 2026-07-05, hypothesis 2): the live LLM extractor is NON-DETERMINISTIC and
BUDGET-COLLAPSES on the epitope lane — the same Keskin paper yields epitopes 2 / 2 / 99 run-to-run (2 of
3 runs drop the whole minimal-epitope layer). That layer lives in a clean, wide supplementary manifest,
one row per predicted minimal epitope. Reading the manifest is a PURE FUNCTION of the sheet, so the
output is byte-identical every run and the collapse cannot happen.

Verified against three gold papers (offline, no quota):
  - Keskin 30568305 "Table S5" (two-row merged header, class I only):  epitopes 99/99, peptides 103/103 EXACT
  - Rojas 37165196 MOESM4 (single-row header, MHC-I + MHC-II columns):  epitopes 451/451, peptides 223/223 EXACT
  - Li    33879241 "Table S2" (BOUNDARY): the minimal epitope is NOT a column — the sheet has only the long
    21-mer; the gold 8-mer is a *derived* NetMHC minimal within it. The loader returns the peptides it can
    source and epitopes=[]; this shape needs a prediction step, not a column read. (Documented, not solved.)

The loader is header-DRIVEN via a synonym registry (not fixed indices), so it survives column reordering
and the 1-row-vs-2-row-merged-header difference across papers of this manifest class. It generalizes the
`tcr_track_adapter.py` pattern to the epitope/immunizing-peptide manifest.

WIRED INTO THE LIVE EXTRACTOR (2026-07-05): `agent_core.load_epitope_manifest_into` merges this loader's
output into the partial record, exposed as the `add_epitope_manifest` MCP tool. Output here is schema-
shaped (validates against schema.MinimalEpitope / schema.ImmunizingPeptide): class-I epitopes always carry
a predicted_affinity slot (lossless unit='unknown' when the source prints no IC50/%rank), and every entity
carries quoted_text/section_ref.

IDS (2026-08-16): immunizing-peptide paper_local_ids are minted by `_imp_id` — a pure function of
the peptide SEQUENCE, collision-proof across containment variants (see that docstring).
"""
from __future__ import annotations

import hashlib
import pathlib
import re

import openpyxl

_PEPTIDE = re.compile(r"^[ACDEFGHIKLMNPQRSTVWY]{7,}$")   # a real amino-acid string, not 'NA' / '-' / blank
_SUB_TOKENS = ("sequence", "affinity", "id")             # markers of a 2-row merged sub-header


def _norm(v):
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v not in (None, "") else ""


def _peptide(v):
    """Return the cell as a clean peptide sequence, or None if it isn't one (drops 'NA', '-', numbers)."""
    if not isinstance(v, str):
        return None
    s = v.strip().upper()
    return s if _PEPTIDE.match(s) else None


def _fmt_hla(raw):
    """'A01:01' -> 'HLA-A*01:01'; leave already-qualified ('HLA-...', 'H-2Kb', 'HLA-DRB1*..') verbatim."""
    if not raw:
        return None
    s = str(raw).strip()
    if s.upper().startswith(("HLA", "H-2", "H2")):
        return s
    m = re.match(r"^([A-Z]+)(\d{2}:\d{2})$", s)
    return f"HLA-{m.group(1)}*{m.group(2)}" if m else s


def _affinity(val, cls, *, require_slot=False):
    """A predicted-affinity Measurement. NetMHCpan nM is a CLASS-I quantity, so a parsed number is only
    emitted as nM for class I. `require_slot` (class-I predicted_affinity) guarantees a lossless
    unit='unknown' Measurement when the source omits the number — the schema requires the slot to exist."""
    if cls == "I" and isinstance(val, (int, float)):
        v = round(float(val), 2)
        return {"value": v, "unit": "nM", "raw": f"{v} nM", "method": "NetMHCpan", "tier": "predicted"}
    if require_slot:  # class-I with no printed IC50/%rank -> lossless slot (Rojas Supp Table 5 shape)
        raw = str(val).strip() if val not in (None, "") else "not reported"
        return {"unit": "unknown", "raw": raw, "method": "NetMHCpan", "tier": "predicted"}
    return None


# Mirror of schema._PROTEIN_CHANGE_PATTERN — protein_change accepts only this shape; anything else
# (e.g. 'p.V30del223,insG', which carries a comma) falls back to the free-form `mutation` field.
_HGVS = re.compile(r"^p\.\(?[A-Za-z0-9_*=>]+\)?$")


def _hgvs(prot):
    s = str(prot).strip() if prot not in (None, "") else ""
    return s if _HGVS.match(s) else None


def _imp_id(seq):
    """`IMP_<first 12 aa>_<length>_<sha1[:6]>` — the immunizing-peptide paper_local_id.

    BUG (live Keskin run, 2026-08-16): the old id was `IMP_{seq[:12]}` alone. Supplementary
    manifests routinely carry CONTAINMENT VARIANTS — the same neoantigen synthesized at two
    lengths ('LARDIPPAVTGKWKLSDLRRYGA' vs '...YGAVPSG'). Those share their first 12 residues, so
    both rows minted the SAME id and `finalize` hard-failed with
    `duplicate paper_local_id in immunizing_peptides` (6 such pairs in Keskin PMID 30568305,
    outputs/tcr_phase5/30568305.log). No allow-flag overrides that check, so the whole paper
    produced no record.

    Stability: this is a PURE FUNCTION OF THE SEQUENCE ALONE — no row order, no counter, no
    randomness, no dependence on which other peptides happen to be in the sheet. The same peptide
    therefore gets the same id in every run and in every paper, which is what the loaders' whole
    byte-identical-reproducibility contract rests on (sha1 is stable across processes; Python's
    `hash()` deliberately is NOT used).

    Uniqueness: the readable `seq[:12]` prefix stays first so curators can still eyeball the id;
    `len(seq)` alone already separates every containment variant seen in real data, and the 6-hex
    sha1 tail of the FULL sequence covers the residual case (same prefix AND same length, differing
    downstream — e.g. a mutant/WT long-peptide pair whose substitution sits past position 12).
    Distinct sequences can only collide on a 24-bit sha1 prefix collision.
    """
    digest = hashlib.sha1(seq.encode("ascii"), usedforsecurity=False).hexdigest()[:6]
    return f"IMP_{seq[:12]}_{len(seq)}_{digest}"


def _is_ii(label):
    return any(t in label for t in ("mhc-ii", "mhc class ii", "mhc ii", "drb", "class ii"))


def _combined_labels(rows, hi):
    """Build one label per column. If the row under the header carries sub-tokens (Keskin's merged
    'Mutated peptide' / 'Sequence' two-row header), fold the group label forward across blank cells and
    append the sub-label; otherwise use the single header row verbatim (Rojas / Li)."""
    grp = rows[hi]
    sub = rows[hi + 1] if hi + 1 < len(rows) else []
    two_row = any(_norm(c) in _SUB_TOKENS for c in sub)
    labels, cur = [], ""
    for i in range(max(len(grp), len(sub) if two_row else 0, len(grp))):
        g = _norm(grp[i]) if i < len(grp) else ""
        if g:
            cur = g
        if two_row:
            s = _norm(sub[i]) if i < len(sub) else ""
            labels.append((f"{cur} {s}".strip(), i))
        else:
            labels.append((g, i))
    return labels, (hi + 2 if two_row else hi + 1)


def _score_header_row(row):
    toks = [_norm(c) for c in row]
    return sum(any(k in t for k in ("allele", "epitope", "peptide", "gene", "neoantigen", "mutation"))
               for t in toks if t)


def _column_map(rows):
    """Locate the header and return (data_start_idx, colmap). colmap keys:
       gene, protein_change, patient, alleles{class:col}, mut_epi{class:col}, wt_epi{class:col},
       long_pep[col,...], mut_aff, wt_aff."""
    hi = max(range(min(len(rows), 12)), key=lambda i: _score_header_row(rows[i]), default=0)
    if _score_header_row(rows[hi]) < 2:
        return len(rows), {}
    labels, data_start = _combined_labels(rows, hi)

    cm = {"alleles": {}, "mut_epi": {}, "wt_epi": {}, "long_pep": []}
    for L, i in labels:
        if not L:
            continue
        cls = "II" if _is_ii(L) else "I"
        is_wt = ("wt" in L.split() or "wild" in L)
        if L == "gene" or L.startswith("gene "):
            cm.setdefault("gene", i)
        elif "ept peptide id" in L or L == "ept id" or L.endswith(" ept id"):
            cm.setdefault("ept_id", i)
        elif "protein change" in L or "substitution" in L or L == "mutation":
            cm.setdefault("protein_change", i)
        elif "patient" in L:
            cm.setdefault("patient", i)
        elif "allele" in L:
            cm["alleles"].setdefault(cls, i)
        elif ("epitope" in L or ("mutated peptide" in L and "sequence" in L)) and ("mutant" in L or "mutated" in L) and not is_wt:
            cm["mut_epi"].setdefault(cls, i)
        elif is_wt and ("epitope" in L or ("wild type peptide" in L and "sequence" in L)):
            cm["wt_epi"].setdefault(cls, i)
        elif ("immunizing peptide" in L and "sequence" in L) or "mutant neoantigen sequence" in L \
                or "mer seq" in L or ("neoantigen" in L and "sequence" in L and not is_wt):
            cm["long_pep"].append(i)
        elif "mutated peptide" in L and "affinity" in L:
            cm.setdefault("mut_aff", i)
        elif "wild type peptide" in L and "affinity" in L:
            cm.setdefault("wt_aff", i)
    return data_start, cm


def _cell(row, i):
    return row[i] if i is not None and i < len(row) else None


def load_epitope_manifest(xlsx_path, sheet=None):
    """Pure function: read the class-I/II epitope manifest -> schema-shaped lanes, deduped by sequence
    (the stable key the ruler scores on), first row wins. Empty lanes if no matching header is found."""
    wb = openpyxl.load_workbook(pathlib.Path(xlsx_path), read_only=True, data_only=True)
    ws = wb[sheet] if sheet else _pick_sheet(wb)
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    data_start, cm = _column_map(rows)
    if not cm:
        return {"epitopes": [], "immunizing_peptides": []}

    epitopes, imms, seen_ep, seen_imm = [], [], set(), set()
    for row in rows[data_start:]:
        gene = _cell(row, cm.get("gene"))
        prot = _cell(row, cm.get("protein_change"))

        # long / immunizing peptide(s) — collected per row, even where an epitope cell is blank
        for pc in cm["long_pep"]:
            seq = _peptide(_cell(row, pc))
            if seq and seq not in seen_imm:
                seen_imm.add(seq)
                imms.append({
                    "paper_local_id": _imp_id(seq),
                    "sequence": seq,
                    "gene_symbol": str(gene).strip() if gene else None,
                    "is_neoantigen": True,
                    "mutation": (str(prot).strip() or None) if prot else None,
                    "protein_change": _hgvs(prot),
                    "quoted_text": f"{str(gene).strip() if gene else ''} immunizing peptide {seq}".strip(),
                    "section_ref": "epitope manifest",
                    "needs_review": False,
                    "provenance": [],
                })

        # long peptides already collected for this row — skip EPTs that stick out past them
        # (Ott COL22A1 LPNEYAFVTT is a 10-mer not contained in the 23-mer IMP).
        row_longs = [_peptide(_cell(row, pc)) for pc in cm["long_pep"]]
        row_longs = [s for s in row_longs if s]
        ept_id_raw = _cell(row, cm.get("ept_id"))
        ept_id = str(ept_id_raw).strip() if ept_id_raw not in (None, "") else ""

        # minimal (short) epitope(s) — one per MHC class that has an explicit mutant-epitope column
        for cls, ec in cm["mut_epi"].items():
            mut = _peptide(_cell(row, ec))
            if not mut or mut in seen_ep:
                continue
            if row_longs and not any(mut in lp for lp in row_longs):
                continue
            seen_ep.add(mut)
            hla = _fmt_hla(_cell(row, cm["alleles"].get(cls)))
            wt = _peptide(_cell(row, cm["wt_epi"].get(cls)))
            q = f"{gene} {mut} ({hla}) predicted MHC-{cls} epitope"
            if ept_id:
                q = f"{q} {ept_id}"
            epitopes.append({
                "paper_local_id": f"EPI_{mut}",
                "sequence": mut,
                "gene_symbol": str(gene).strip() if gene else None,
                "is_neoantigen": True,
                "hla_allele": hla,
                "mhc_class": cls,
                "wild_type_sequence": wt,
                "predicted_affinity": _affinity(_cell(row, cm.get("mut_aff")), cls, require_slot=(cls == "I")),
                "wild_type_affinity": _affinity(_cell(row, cm.get("wt_aff")), cls),
                "section_ref": "epitope manifest",
                "quoted_text": q,
                "needs_review": False,
                "provenance": [],
            })
    return {"epitopes": epitopes, "immunizing_peptides": imms}


def _pick_sheet(wb):
    """Pick the sheet whose best header row scores highest on manifest tokens; ties -> first sheet."""
    best, best_score = wb.worksheets[0], -1
    for ws in wb.worksheets:
        head = list(ws.iter_rows(min_row=1, max_row=6, values_only=True))
        score = max((_score_header_row(r) for r in head), default=0)
        if score > best_score:
            best, best_score = ws, score
    return best
