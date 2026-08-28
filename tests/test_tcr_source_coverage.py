"""SOURCE UNDER-COVERAGE gate (2026-08-24): census + load ledger + `_tcr_source_coverage_gap`.

Root cause it locks down: every other finalize gate checks a record against ITSELF, so a record that
read a FRACTION of what the supplement offered finalized clean -- Hu 33479501's keep-file asserts
tcr_seq_status='both' with 3 TCR-seq methods and ZERO clonotypes beside a supplement holding 10
loadable tet-spec track sheets (269 clonotypes / 654 observations), and a 3-run replicate routed
11 / 6 / 1 sheets to the loaders on identical input with all three finalizing clean.

Covers: (A) SHAPE-not-name classification, incl. the two look-alikes that must NOT become required
(the unsorted bulk repertoire sheet with the same identity column, and the single-chain bulk-RNA
sheet with the same column NAMES as the paired reactive shape); (B) the curation OUT-OF-SCOPE rule
that keeps Keskin's non-antigen-sorted 'Tumor-associated' table from being coerced into records;
(C) the load ledger the loaders write; (D) the gate -- every SILENT degradation path, the block, the
named-sheet message, and the HARD override that routes to needs_review.
"""
import json
import pathlib

import openpyxl
import pytest

import agent_core as ac
from vaxtract import tcr_source_census as census

META = json.dumps({"pmid": "33479501", "title": "t", "journal": "j", "year": 2021,
                   "cohort_size": 1, "indication_summary": "melanoma"})
PROV = {"quoted_text": "q", "section_ref": "Suppl 8"}

# The Hu track shape: title row, then the timepoint-label header (row index 1), then data.
TRACK_HEADER = ["CDR3β", "1=yes", "Pre-Vax", "Week 4", "Long-term"]
TRACK_ROWS = [["CASSLGNQPQHF", 1, 0.01, 0.05, None],
              ["CASSDISYEQYF", 1, None, None, 0.007]]
# The Keskin paired single-cell reactive shape: one row per CELL.
REACTIVE_HEADER = ["Cell designation", "TRAV", "TRAJ", "Alpha CDR3 DNA seq",
                   "Alpha CDR3 amino acid seq", "TRBV", "TRBJ", "TRBC", "Beta CDR3 DNA seq",
                   "Beta CDR3 amino acid seq", "Clone", "Clonotype"]
REACTIVE_ROWS = [["c1", "TRAV1", "TRAJ2", "gat", "CAVRDGGATNKLIF", "TRBV5-1", "TRBJ2-1", "TRBC2",
                  "gat", "CASSLDRNNEQFF", "1", "1"],
                 ["c2", "TRAV1", "TRAJ2", "gat", "CAVRDGGATNKLIF", "TRBV5-1", "TRBJ2-1", "TRBC2",
                  "gat", "CASSLDRNNEQFF", "1", "1"]]


def _book(tmp_path, name, sheets):
    """sheets = [(sheet_title, title_cell, header, rows)] -> one .xlsx in a fresh paper dir."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, title_cell, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append([title_cell])
        ws.append(header)
        for r in rows:
            ws.append(r)
    p = pathlib.Path(tmp_path) / name
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p


def _track_book(tmp_path, sheet="Suppl 8 Pt2 Tet-spec TCRb track"):
    return _book(tmp_path, "supp3.xlsx", [(sheet, "Pt2 tetramer-specific TCRb", TRACK_HEADER,
                                           TRACK_ROWS)])


def _record(tmp_path, paper_dir=None, name="o.json"):
    """A record that can reach the finalize gates: one patient (the loaders' 'Pt2') + one peptide."""
    out = str(pathlib.Path(tmp_path) / name)
    if paper_dir is not None:
        assert ac.set_source_dir(out, str(paper_dir))[0]
    assert ac.init_partial(out, META)[0]
    ac.append_section(out, "patients", json.dumps([{"paper_local_id": "Pt2", "indication": "melanoma",
                                                    "n_peptides_synthesized": 1,
                                                    "n_peptides_immunogenic": 1, **PROV}]))
    ac.append_section(out, "immunizing_peptides", json.dumps(
        [{"paper_local_id": "IMP1", "sequence": "PQVDGEIPLHRSDRVKVLSIGEGGF", "gene_symbol": "SHANK2",
          "is_neoantigen": True, **PROV}]))
    return out


# ------------------------------------------------------------------ A. shape, not name

def test_track_sheet_is_required_by_shape_not_name(tmp_path):
    # A sheet called 'Sheet1' with the track SHAPE is still required: Hu's own names are
    # inconsistent ('Suppl 8 Pt 5Tet-spec TCRa track' has no space), so names cannot be the signal.
    _track_book(tmp_path, sheet="Sheet1")
    c = census.census_tcr_sheets(tmp_path)
    assert [(s.sheet, s.shape, s.loader) for s in c.required] == [("Sheet1", "track", "add_tcr_track")]
    assert c.out_of_scope == []

def test_locus_is_inferred_from_the_identity_header(tmp_path):
    _book(tmp_path, "s.xlsx", [("a", "t", ["TCR⍺", "1=yes", "Pre-Vax"], [["CAVRDGGATNKLIF", 1, 0.1]])])
    assert "TRA identity column" in census.census_tcr_sheets(tmp_path).required[0].reason

def test_unnamed_chain_is_not_required(tmp_path):
    # `tcr_track_adapter` REFUSES to guess the locus, so a sheet it would refuse is not required.
    _book(tmp_path, "s.xlsx", [("a", "t", ["Sequence", "1=yes", "Pre-Vax"], [["CASSLGNQPQHF", 1, .1]])])
    c = census.census_tcr_sheets(tmp_path)
    assert c.required == [] and ("s.xlsx", "a") in c.not_loadable

def test_bulk_repertoire_lookalike_is_not_required(tmp_path):
    # Hu's 'Suppl 8 Pt 2 All CDR3b in bulk' has the SAME identity column as the track sheet next to
    # it; only the paper's specificity flag in col 1 separates them, and the loader needs that flag.
    _book(tmp_path, "s.xlsx", [("bulk", "t", ["CDR3β", "UMI per clonotype", "Frequency", "%"],
                                [["CASSLGNQPQHF", 704, 0.011, 1.19]])])
    assert census.census_tcr_sheets(tmp_path).required == []

def test_reactive_sheet_is_required(tmp_path):
    _book(tmp_path, "s.xlsx", [("(b)Pt7 Neoantigen-reactive CD4+", "Pt7 PoolC reactive CD4+ T cells",
                                REACTIVE_HEADER, REACTIVE_ROWS)])
    s, = census.census_tcr_sheets(tmp_path).required
    assert (s.shape, s.loader) == ("reactive", "add_reactive_tcr")

def test_single_chain_bulk_rna_lookalike_is_not_required(tmp_path):
    # Keskin MOESM8 ('CDR3 sequences from analysis of bulk RNA') carries TRAV/'Alpha CDR3 amino acid
    # seq' -- the same column NAMES as the reactive shape -- but one row per CLONOTYPE, not per CELL.
    # Feeding it to the reactive adapter would mint thousands of unsorted pseudo-clonotypes.
    _book(tmp_path, "s.xlsx", [("a) Blood Pre-Vac Alpha", "Patient 8 CDR3 sequences from bulk RNA",
                                ["TRAV", "TRAJ", "Alpha CDR3 amino acid seq", "UMI_count"],
                                [["TRAV1", "TRAJ2", "CAVRDGGATNKLIF", 12]])])
    assert census.census_tcr_sheets(tmp_path).required == []

def test_non_tcr_sheet_is_not_loadable(tmp_path):
    _book(tmp_path, "s.xlsx", [("QC", "QC metrics", ["Patient ID", "Coverage"], [["Pt1", 40]])])
    c = census.census_tcr_sheets(tmp_path)
    assert c.required == [] and c.not_loadable == [("s.xlsx", "QC")]

def test_header_with_no_data_row_is_not_required(tmp_path):
    # An empty table offers nothing; demanding a load that yields nothing would be unfixable.
    _book(tmp_path, "s.xlsx", [("a", "t", TRACK_HEADER, [])])
    assert census.census_tcr_sheets(tmp_path).required == []


# ------------------------------------------------------------------ B. the curation out-of-scope rule

def test_tumour_associated_sheet_is_out_of_scope_not_required(tmp_path):
    # POLICY (awaiting maintainer sign-off): Keskin's '(a)Pt7 Tumor-associated Tcells' is a valid
    # reactive SHAPE but is not antigen-sorted. If it were required, the gate would coerce that
    # 228-row contamination into every Keskin record.
    _book(tmp_path, "s.xlsx", [("(a)Pt7 Tumor-associated Tcells", "Pt7 Tumor-associated T cells",
                                REACTIVE_HEADER, REACTIVE_ROWS)])
    c = census.census_tcr_sheets(tmp_path)
    assert c.required == []
    assert [(s.sheet, s.scope) for s in c.out_of_scope] == \
        [("(a)Pt7 Tumor-associated Tcells", census.OUT_OF_SCOPE)]
    assert "not antigen-sorted" in c.out_of_scope[0].reason

def test_out_of_scope_matches_the_title_cell_too(tmp_path):
    # The exclusion cue lives in the sheet NAME or the table's own title row -- the two places a
    # supplement says what its cells were sorted on.
    _book(tmp_path, "s.xlsx", [("Table 12", "Tumour-infiltrating T cells, patient 7",
                                REACTIVE_HEADER, REACTIVE_ROWS)])
    assert census.census_tcr_sheets(tmp_path).required == []

def test_out_of_scope_rule_is_data_not_a_hardcoded_sheet(tmp_path):
    # The rule must GENERALISE -- it is a documented pattern over what a table says it holds, not a
    # (file, sheet) allow-list. The same content in a differently named file/sheet is still excluded.
    assert [name for name, _, _ in census.OUT_OF_SCOPE_RULES] == ["tumour_associated", "unsorted"]
    _book(tmp_path, "some_other_paper_S4.xlsx",
          [("T cell receptors", "Tumor-associated T cells, subject 12", REACTIVE_HEADER,
            REACTIVE_ROWS)])
    c = census.census_tcr_sheets(tmp_path)
    assert c.required == [] and len(c.out_of_scope) == 1


# ------------------------------------------------------------------ C. the load ledger

def test_track_loader_writes_the_ledger(tmp_path):
    paper = tmp_path / "paper"
    path = _track_book(paper)
    out = _record(tmp_path, paper)
    assert ac.load_tcr_track_into(out, str(path), "Suppl 8 Pt2 Tet-spec TCRb track", "Pt2")[0]
    entry, = ac.read_load_ledger(out)
    assert entry["file"] == "supp3.xlsx"            # BASENAME, so path spelling never mismatches
    assert entry["sheet"] == "Suppl 8 Pt2 Tet-spec TCRb track"
    assert (entry["loader"], entry["n_loaded"], entry["n_available"]) == ("add_tcr_track", 2, 2)

def test_reload_replaces_the_entry_and_counts_attempts(tmp_path):
    paper = tmp_path / "paper"
    path = _track_book(paper)
    out = _record(tmp_path, paper)
    sheet = "Suppl 8 Pt2 Tet-spec TCRb track"
    ac.load_tcr_track_into(out, str(path), sheet, "Pt2")
    ac.load_tcr_track_into(out, str(path), sheet, "Pt2")     # idempotent re-load: 0 fresh rows
    entry, = ac.read_load_ledger(out)
    assert entry["attempts"] == 2 and entry["n_loaded"] == 0 and entry["n_available"] == 2

def test_empty_answer_is_still_ledgered(tmp_path):
    # A routed sheet the adapter reads as EMPTY is an answer about that sheet; if it did not ledger,
    # the gate would demand it forever and no tool call could satisfy the block.
    paper = tmp_path / "paper"
    path = _book(paper, "s.xlsx", [("a", "t", TRACK_HEADER, [["CASSLGNQPQHF", 0, 0.1, None, None]])])
    out = _record(tmp_path, paper)
    ok, msg = ac.load_tcr_track_into(out, str(path), "a", "Pt2")
    assert not ok and "no tetramer-specific" in msg
    assert ac.read_load_ledger(out) == [{"file": "s.xlsx", "sheet": "a", "loader": "add_tcr_track",
                                         "n_loaded": 0, "n_available": 0, "attempts": 1}]

def test_reactive_loader_writes_the_ledger(tmp_path):
    paper = tmp_path / "paper"
    path = _book(paper, "m7.xlsx", [("(b)Pt8 mutSVEP1-react Tcells", "Pt8 mutantSVEP1 reactive",
                                     REACTIVE_HEADER, REACTIVE_ROWS)])
    out = _record(tmp_path, paper)
    assert ac.load_reactive_tcr_into(out, str(path), "(b)Pt8 mutSVEP1-react Tcells", "Pt8",
                                     "mutant SVEP1")[0]
    entry, = ac.read_load_ledger(out)
    assert (entry["file"], entry["loader"], entry["n_available"]) == ("m7.xlsx", "add_reactive_tcr", 1)

def test_init_record_clears_a_stale_ledger(tmp_path):
    paper = tmp_path / "paper"
    path = _track_book(paper)
    out = _record(tmp_path, paper)
    ac.load_tcr_track_into(out, str(path), "Suppl 8 Pt2 Tet-spec TCRb track", "Pt2")
    pathlib.Path(out + ".locked.json").unlink()      # allow the re-init this test is about
    assert ac.init_partial(out, META)[0]
    assert ac.read_load_ledger(out) == []


# ------------------------------------------------------------------ D. the gate

def _paper(out, **kw):
    rec = json.loads(pathlib.Path(out + ".partial.json").read_text())
    rec.update(kw)
    return ac.ExtractedPaper(**rec)


def test_gap_fires_and_names_the_sheet_and_its_loader(tmp_path):
    paper = tmp_path / "paper"
    _track_book(paper)
    out = _record(tmp_path, paper)
    gap = ac._tcr_source_coverage_gap(_paper(out), out)
    assert "1 loadable TCR sheet(s) but this record loaded 0" in gap
    assert "supp3.xlsx [Suppl 8 Pt2 Tet-spec TCRb track] -> add_tcr_track" in gap

def test_gap_silent_once_the_sheet_is_loaded(tmp_path):
    paper = tmp_path / "paper"
    path = _track_book(paper)
    out = _record(tmp_path, paper)
    ac.load_tcr_track_into(out, str(path), "Suppl 8 Pt2 Tet-spec TCRb track", "Pt2")
    assert ac._tcr_source_coverage_gap(_paper(out), out) == ""

def test_gap_silent_for_an_honest_zero(tmp_path):
    # Rojas 37165196: status='both', zero clonotypes, and NO clonotype-shaped sheet in the supps.
    # A false positive here would be the serious regression.
    paper = tmp_path / "paper"
    _book(paper, "s.xlsx", [("Figure 2A", "Figure 2A, Bar graph", ["", "TCRVβ +"], [["Pt1", 3]])])
    out = _record(tmp_path, paper)
    assert ac._tcr_source_coverage_gap(_paper(out, tcr_seq_status="both"), out) == ""

def test_gap_silent_without_a_source_pointer(tmp_path):
    # e.g. `refinalize`, or any record finalized outside the extraction runtime.
    paper = tmp_path / "paper"
    _track_book(paper)
    out = _record(tmp_path, paper_dir=None)
    assert ac._tcr_source_coverage_gap(_paper(out), out) == ""

def test_gap_silent_when_the_source_dir_is_gone(tmp_path):
    paper = tmp_path / "paper"
    path = _track_book(paper)
    out = _record(tmp_path, paper)
    path.unlink(); paper.rmdir()
    assert ac._tcr_source_coverage_gap(_paper(out), out) == ""

def test_gap_silent_when_the_pointer_names_another_paper(tmp_path):
    # A reused out_path must not import another paper's supplements as this record's requirements.
    paper = tmp_path / "paper"
    _track_book(paper)
    out = _record(tmp_path, paper)
    meta = json.loads(pathlib.Path(out + ".source.json").read_text())
    assert meta["pmid"] == "33479501"                       # stamped by init_record
    meta["pmid"] = "99999999"
    pathlib.Path(out + ".source.json").write_text(json.dumps(meta))
    assert ac._tcr_source_coverage_gap(_paper(out), out) == ""

def test_gap_silent_on_an_unreadable_workbook(tmp_path):
    paper = tmp_path / "paper"
    paper.mkdir()
    (paper / "broken.xlsx").write_bytes(b"not a zip")
    out = _record(tmp_path, paper)
    assert ac._tcr_source_coverage_gap(_paper(out), out) == ""
    assert census.census_tcr_sheets(paper).files_failed == [str(paper / "broken.xlsx")]

def test_out_of_scope_sheet_never_fires_the_gate(tmp_path):
    paper = tmp_path / "paper"
    _book(paper, "m10.xlsx", [("(a)Pt7 Tumor-associated Tcells", "Pt7 Tumor-associated T cells",
                               REACTIVE_HEADER, REACTIVE_ROWS)])
    out = _record(tmp_path, paper)
    assert ac._tcr_source_coverage_gap(_paper(out), out) == ""


# ------------------------------------------------------------------ E. finalize wiring

def test_finalize_blocks_then_the_override_routes_to_needs_review(tmp_path):
    paper = tmp_path / "paper"
    _track_book(paper)
    out = _record(tmp_path, paper)
    ok, msg = ac.finalize_partial(out)
    assert not ok and "TCR source coverage" in msg and "add_tcr_track" in msg
    ok2, msg2 = ac.finalize_partial(out, allow_unloaded_tcr_sheets=True)
    assert ok2
    rec = json.loads(pathlib.Path(out).read_text())
    assert rec["finalize_overrides_used"] == ["allow_unloaded_tcr_sheets"]

def test_loading_the_sheet_clears_the_block(tmp_path):
    paper = tmp_path / "paper"
    path = _track_book(paper)
    out = _record(tmp_path, paper)
    ac.load_tcr_track_into(out, str(path), "Suppl 8 Pt2 Tet-spec TCRb track", "Pt2")
    ok, msg = ac.finalize_partial(out, allow_missing_tcr_flags=True, allow_missing_tcr_gateway=True)
    assert ok, msg
    assert "allow_unloaded_tcr_sheets" not in \
        json.loads(pathlib.Path(out).read_text())["finalize_overrides_used"]

def test_override_is_hard_so_the_scale_lane_reviews_it():
    assert "allow_unloaded_tcr_sheets" in ac.HARD_OVERRIDES
    assert "allow_unloaded_tcr_sheets" not in ac.SOFT_OVERRIDES

def test_refinalize_clears_the_sidecars_and_stays_silent(tmp_path):
    paper = tmp_path / "paper"
    _track_book(paper)
    out = _record(tmp_path, paper)
    ok, _ = ac.finalize_partial(out, allow_unloaded_tcr_sheets=True)
    assert ok
    # A stale pointer/ledger at the destination must not be read as the replayed record's coverage.
    dest = str(tmp_path / "replay.json")
    ac.set_source_dir(dest, str(paper))
    ok2, msg2 = ac.refinalize(out, dest)
    assert ok2, msg2
    assert not pathlib.Path(dest + ".source.json").exists()
    assert json.loads(pathlib.Path(dest).read_text())["finalize_overrides_used"] == []


# ------------------------------------------------------------------ F. the fast shared-string reader

@pytest.mark.parametrize("chunk", [1 << 22, 64, 7])
def test_shared_strings_resolve_across_chunk_boundaries(tmp_path, monkeypatch, chunk):
    """The census never opens the workbook with openpyxl (Hu's 159MB / 4.2M-entry shared-string table
    costs ~13s to parse before a single cell is read); it scans the table forward and stops at the
    largest index needed. Tiny chunk sizes force every straddle path in that scan."""
    monkeypatch.setattr(census, "_SS_CHUNK", chunk)
    paper = tmp_path / "paper"
    filler = [[f"padding string {i}", i] for i in range(200)]   # push the header off index 0
    _book(paper, "s.xlsx", [("filler", "titl", ["a", "b"], filler),
                            ("track", "Pt2 tetramer-specific TCRb", TRACK_HEADER, TRACK_ROWS)])
    c = census.census_tcr_sheets(paper)
    assert [s.sheet for s in c.required] == ["track"]


# ---------------------------------------------------------------------------------------------
# add_table now ledgers too. A generic bulk-load must NEVER be mistaken for the deterministic TCR
# load this gate demands -- that would turn the gate silently OFF, which is its worst failure mode.

def test_add_table_load_does_not_satisfy_the_tcr_gate(tmp_path):
    paper = tmp_path / "paper"
    path = _track_book(paper)
    out = _record(tmp_path, paper)
    ac._append_load_ledger(out, path=str(path), sheet="Suppl 8 Pt2 Tet-spec TCRb track",
                           loader="add_table", n_loaded=99, n_available=99)
    gap = ac._tcr_source_coverage_gap(_paper(out), out)
    assert "this record loaded 0" in gap, "add_table must not count as a TCR load"


# ---------------------------------------------------------------------------------------------
# READ under-coverage gate: what the agent was SHOWN vs what it kept.

@pytest.fixture(autouse=True)
def _clean_read_observations():
    ac.reset_read_observations()
    yield
    ac.reset_read_observations()


def test_read_gap_fires_when_a_filtered_read_is_barely_loaded(tmp_path):
    out = str(tmp_path / "r.json")
    ac.init_partial(out, META)
    ac._observe_filtered_read("/x/supp3.xlsx", "Dataset 4b", matched=121, total=375,
                              row_filter={"col_idx": 10, "equals": 1})
    gap = ac._read_under_coverage_gap(out)
    assert "supp3.xlsx [Dataset 4b]" in gap and "matched 121/375" in gap
    assert "loaders took 0" in gap


def test_read_gap_silent_once_a_loader_takes_the_rows(tmp_path):
    out = str(tmp_path / "r.json")
    ac.init_partial(out, META)
    ac._observe_filtered_read("/x/supp3.xlsx", "Dataset 4b", matched=121, total=375,
                              row_filter={"col_idx": 10, "equals": 1})
    # per_cell fan-out routinely returns MORE rows than the filter matched (121 -> 279)
    ac._append_load_ledger(out, path="/x/supp3.xlsx", sheet="Dataset 4b", loader="add_table",
                           n_loaded=279, n_available=375)
    assert ac._read_under_coverage_gap(out) == ""


def test_read_gap_silent_for_a_small_inspection_filter(tmp_path):
    out = str(tmp_path / "r.json")
    ac.init_partial(out, META)
    ac._observe_filtered_read("/x/s.xlsx", "S1", matched=3, total=400, row_filter={"col": "g"})
    assert ac._read_under_coverage_gap(out) == ""


def test_read_gap_silent_when_half_the_shown_rows_are_loaded(tmp_path):
    out = str(tmp_path / "r.json")
    ac.init_partial(out, META)
    ac._observe_filtered_read("/x/s.xlsx", "S1", matched=100, total=400, row_filter={"col": "g"})
    ac._append_load_ledger(out, path="/x/s.xlsx", sheet="S1", loader="add_table",
                           n_loaded=50, n_available=400)
    assert ac._read_under_coverage_gap(out) == ""


def test_read_gap_silent_with_no_observations(tmp_path):
    # every record finalized outside a live run (refinalize, the committed artifacts) -> inert
    out = str(tmp_path / "r.json")
    ac.init_partial(out, META)
    assert ac._read_under_coverage_gap(out) == ""


def test_repeat_read_keeps_the_largest_match(tmp_path):
    ac._observe_filtered_read("/x/s.xlsx", "S1", matched=121, total=375, row_filter={"col": "a"})
    ac._observe_filtered_read("/x/s.xlsx", "S1", matched=4, total=375, row_filter={"col": "b"})
    assert ac._READ_OBSERVATIONS[("s.xlsx", "S1")]["matched"] == 121


def test_init_record_clears_stale_read_observations(tmp_path):
    ac._observe_filtered_read("/x/s.xlsx", "S1", matched=121, total=375, row_filter={"col": "a"})
    ac.init_partial(str(tmp_path / "r.json"), META)
    assert ac._READ_OBSERVATIONS == {}


def test_unfiltered_read_is_not_observed(tmp_path):
    """An unfiltered preview reports the WHOLE sheet (375 rows of Hu 4b, 121 of them positive);
    gating on that total would false-block every honest load."""
    book = tmp_path / "s.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Patient", "Flag"])
    for i in range(30):
        ws.append([f"Pt{i}", 1])
    wb.save(book)
    ac.read_table_rows(str(book))
    assert ac._READ_OBSERVATIONS == {}
    ac.read_table_rows(str(book), row_filter={"col": "Flag", "equals": 1})
    assert ac._READ_OBSERVATIONS[(book.name.lower(), "Sheet")]["matched"] == 30


def test_read_gap_blocks_finalize_then_the_override_routes_to_needs_review(tmp_path):
    out = str(tmp_path / "r.json")
    ac.init_partial(out, META)
    ac._observe_filtered_read("/x/supp3.xlsx", "Dataset 4b", matched=121, total=375,
                              row_filter={"col_idx": 10, "equals": 1})
    ok, msg = ac.finalize_partial(out)
    assert not ok
    assert "source under-coverage" in msg and "Dataset 4b" in msg
    assert "per_cell" in msg                      # points at the tool that fixes it
    ok2, _ = ac.finalize_partial(out, allow_underloaded_read_sheets=True)
    assert ok2
    rec = json.loads(pathlib.Path(out).read_text())
    assert "allow_underloaded_read_sheets" in rec["finalize_overrides_used"]


def test_the_read_override_is_hard_so_it_routes_to_needs_review():
    assert "allow_underloaded_read_sheets" in ac.HARD_OVERRIDES
    assert not ac.overrides_are_soft_only(["allow_underloaded_read_sheets"])


def test_the_read_override_is_exposed_on_the_finalize_tool():
    from vaxtract import tool_registry as tr
    assert "allow_underloaded_read_sheets" in tr.finalize_override_names()
    schema = tr.finalize_input_schema()
    desc = schema["properties"]["allow_underloaded_read_sheets"]["description"]
    assert "override flag" not in desc, "override needs a real description, not the fallback"


# ---------------------------------------------------------------------------------------------
# PR-3: REACTIVITY-MATRIX census + gate. Asks the SOURCE, not the agent's reads -- so it catches a
# sheet the agent never opened. On Hu 33479501 that is literally true: the census finds a third
# reactivity sheet ('Supp11b NetMHCpan2.4 CD4 T cell', 11 positive calls) that NONE of the five
# committed runs cited, because none of them read it.

# One row per assay peptide; two timepoints x two stimulations scored 1 / 0 / n.d.
REACT_TITLE = "Supplementary Dataset 4b. CD4+ T cell reactivity by IFN-g ELISPOT"
REACT_HEADER = ["Patient ID", "Gene", "Assay peptide", "W16 ex vivo", "W16 pre-stim",
                "Yr3 ex vivo", "Yr3 pre-stim"]


def _react_rows(n=12):
    """Two scored assay columns that VARY -- a column that never varies is a constant, not a call
    (see the normalisation-constant test below) -- plus an all-negative and an all-unknown column.
    Positives: 6 in 'W16 ex vivo' + 11 in 'W16 pre-stim' = 17."""
    out = []
    for i in range(n):
        ex_vivo = 1 if i % 2 == 0 else 0
        pre_stim = 1 if i < n - 1 else 0
        out.append([str(i % 6 + 1), f"GENE{i}", f"PEPTIDESEQ{i}", ex_vivo, pre_stim, 0, "n.d."])
    return out


def _reactivity_book(tmp_path, header=REACT_HEADER, rows=None, title=REACT_TITLE,
                     sheet="Suppl Dataset 4b. CD4 T cells"):
    return _book(tmp_path, "supp4.xlsx", [(sheet, title, header, rows or _react_rows())])


def test_census_finds_a_reactivity_matrix_and_counts_its_positive_cells(tmp_path):
    paper = tmp_path / "paper"
    _reactivity_book(paper)
    cen = census.census_tcr_sheets(paper)
    assert len(cen.reactivity) == 1
    s = cen.reactivity[0]
    assert s.shape == "reactivity" and s.n_expected == 17        # 6 + 11 positive calls
    assert "scored assay column" in s.reason
    assert cen.required == []                                    # NOT a TCR sheet


def test_mutated_peptide_column_is_the_assay_peptide(tmp_path):
    """Hu 4a labels the 9-mer 'Mutated peptide Sequence', not 'Assay peptide'."""
    paper = tmp_path / "paper"
    imp = "PQVDGEIPLHRSDRVKVLSIGEGGF"
    header = ["Patient ID", "Immunizing peptide Sequence", "Mutated peptide Sequence",
              "W16 peptide pulsed"]
    rows = []
    for i in range(12):
        nine = imp[i:i + 9]
        rows.append([str(1), imp, nine, 1 if i < 11 else 0])
    _reactivity_book(paper, header=header, rows=rows,
                     title="CD8+ T cell reactivity by IFN-g ELISPOT",
                     sheet="Suppl Dataset 4a. CD8+ T cells")
    s = census.census_tcr_sheets(paper, with_structure=True).reactivity[0]
    assert s.imp_col is not None and s.assay_pep_col is not None
    assert s.imp_col != s.assay_pep_col
    assert s.n_expected == 11


def test_reactivity_sheet_does_not_become_a_required_tcr_sheet(tmp_path):
    """The TCR gate demands a TCR loader; coercing a reactivity matrix into it would emit a wrong
    block message telling the agent to call add_tcr_track on an immunogenicity table."""
    paper = tmp_path / "paper"
    _reactivity_book(paper)
    out = _record(tmp_path, paper)
    assert ac._tcr_source_coverage_gap(_paper(out), out) == ""


def test_an_id_column_of_small_integers_is_not_a_scored_column(tmp_path):
    """The killer false positive: a patient/rank column holding 1,2,3... must not read as an assay
    column. Only a column whose WHOLE domain is the scored vocabulary qualifies."""
    paper = tmp_path / "paper"
    # col 3 holds 1..12 (an id/rank column); col 4 is a real scored column, 11 positives of 12
    rows = [[str(i), f"G{i}", f"SEQ{i}", i, 1 if i < 12 else 0, 0, "n.d."] for i in range(1, 13)]
    _reactivity_book(paper, rows=rows)
    s = census.census_tcr_sheets(paper).reactivity[0]
    assert s.n_expected == 11, "only the scored column should count, not the id column"


def test_a_sheet_without_a_reactivity_header_is_not_censused(tmp_path):
    paper = tmp_path / "paper"
    _reactivity_book(paper, title="Supplementary Dataset 2. Somatic mutations",
                     header=["Patient", "Gene", "Peptide", "Flag A", "Flag B", "Flag C", "Flag D"])
    assert census.census_tcr_sheets(paper).reactivity == []


def test_a_reactivity_sheet_with_too_few_positives_is_not_gated(tmp_path):
    paper = tmp_path / "paper"
    rows = [[str(i), f"G{i}", f"SEQ{i}", 0, 0, 0, "n.d."] for i in range(12)]
    rows[0][3] = 1          # one positive in an otherwise negative column -> below the floor
    _reactivity_book(paper, rows=rows)
    assert census.census_tcr_sheets(paper).reactivity == []


def test_a_track_sheet_is_never_read_as_a_reactivity_matrix(tmp_path):
    """A track sheet's specificity column is all 1s -- it would qualify as a 'scored column' if the
    census ever offered it to phase 2. It must stay claimed by the TCR classifier."""
    paper = tmp_path / "paper"
    _track_book(paper)
    cen = census.census_tcr_sheets(paper)
    assert len(cen.required) == 1 and cen.reactivity == []


def test_reactivity_gap_fires_then_clears_once_the_cells_are_loaded(tmp_path):
    paper = tmp_path / "paper"
    path = _reactivity_book(paper)
    out = _record(tmp_path, paper)
    gap = ac._reactivity_source_coverage_gap(_paper(out), out)
    assert "17 positive assay call(s), loaded 0" in gap
    # a per_cell load produces exactly one row per positive cell
    ac._append_load_ledger(out, path=str(path), sheet="Suppl Dataset 4b. CD4 T cells",
                           loader="add_table", n_loaded=17, n_available=12)
    assert ac._reactivity_source_coverage_gap(_paper(out), out) == ""


def test_reactivity_gap_fires_on_a_partial_row_grain_load(tmp_path):
    """run2's failure: add_table pointed at the sheet but with a plain row mapping, so it reached
    ONE assay column (6 of 17 calls). A third of the sheet, silently."""
    paper = tmp_path / "paper"
    path = _reactivity_book(paper)
    out = _record(tmp_path, paper)
    ac._append_load_ledger(out, path=str(path), sheet="Suppl Dataset 4b. CD4 T cells",
                           loader="add_table", n_loaded=6, n_available=12)
    assert "loaded 6" in ac._reactivity_source_coverage_gap(_paper(out), out)


def test_reactivity_gap_silent_without_a_source_pointer(tmp_path):
    out = str(tmp_path / "r.json")
    ac.init_partial(out, META)
    assert ac._reactivity_source_coverage_gap(_paper(out), out) == ""


def test_reactivity_gap_blocks_finalize_then_the_override_routes_to_needs_review(tmp_path):
    paper = tmp_path / "paper"
    _reactivity_book(paper)
    out = _record(tmp_path, paper)
    ok, msg = ac.finalize_partial(out)
    assert not ok
    assert "evidence source coverage" in msg and "per_cell" in msg
    ok2, _ = ac.finalize_partial(out, allow_underloaded_reactivity_sheets=True,
                                 allow_missing_tcr_flags=True, allow_missing_tcr_gateway=True)
    assert ok2, "override should clear the reactivity gate"
    rec = json.loads(pathlib.Path(out).read_text())
    assert "allow_underloaded_reactivity_sheets" in rec["finalize_overrides_used"]


def test_the_reactivity_override_is_hard_and_documented():
    from vaxtract import tool_registry as tr
    assert "allow_underloaded_reactivity_sheets" in ac.HARD_OVERRIDES
    assert "allow_underloaded_reactivity_sheets" in tr.finalize_override_names()
    desc = tr.finalize_input_schema()["properties"]["allow_underloaded_reactivity_sheets"]["description"]
    assert "override flag" not in desc


def test_a_normalisation_constant_column_is_not_a_scored_column(tmp_path):
    """PMID 40790272 'Ext. Fig 6a': a 'Baseline FC' column that is 1.0 on EVERY row, sitting beside
    a 'Max response FC' header. It passed the whole-domain-scored and has-a-positive conditions and
    the census called it a 10-cell reactivity matrix. A column that never varies discriminates
    nothing, so a scored column must also hold at least one non-positive call."""
    paper = tmp_path / "paper"
    rows = [[str(i), f"G{i}", f"SEQ{i}", 1, 1, 1, 1] for i in range(12)]   # every call positive
    _reactivity_book(paper, rows=rows,
                     header=["Patient ID", "Gene", "Peptide", "Baseline FC", "b", "c", "d"],
                     title="Max response FC by patient")
    assert census.census_tcr_sheets(paper).reactivity == []


def test_a_column_with_positives_and_negatives_still_counts(tmp_path):
    paper = tmp_path / "paper"
    rows = [[str(i), f"G{i}", f"SEQ{i}", i % 2, 0, 0, "n.d."] for i in range(24)]
    _reactivity_book(paper, rows=rows)
    assert census.census_tcr_sheets(paper).reactivity[0].n_expected == 12


def test_structure_is_opt_in_and_never_changes_the_count(tmp_path):
    """The finalize gate needs only the positive-cell count; deriving the column STRUCTURE means
    resolving the sheet's peptide/header strings, ~2s per workbook on Hu. Opt-in keeps the gate at
    its old cost — but skipping it must not change what the gate counts, or the two callers would
    be measuring different things again."""
    paper = tmp_path / "paper"
    _reactivity_book(paper)
    gate = census.census_tcr_sheets(paper).reactivity
    ruler = census.census_tcr_sheets(paper, with_structure=True).reactivity
    assert [s.n_expected for s in gate] == [s.n_expected for s in ruler]
    assert [s.assay_cols for s in gate] == [s.assay_cols for s in ruler]
    # structure fields are populated only on the ruler path (this fixture has no peptide column,
    # so `patient_col` is the field that shows the difference)
    assert gate[0].patient_col is None and ruler[0].patient_col == 0


def test_the_block_message_hands_over_the_columns(tmp_path):
    """A per_cell mapping is authored fresh every run, so an INFERRED one is the step the run-to-run
    swing moved to once the loader existed: on Hu 33479501 with identical code, one run wrote sheet
    4b exactly (279 calls) and the next wrote `equals: 0` rules on 4a as well (36 rows against 18
    calls). The census already knows the columns, so the block message states them."""
    paper = tmp_path / "paper"
    _reactivity_book(paper)
    out = _record(tmp_path, paper)
    msg = ac._reactivity_source_coverage_gap(_paper(out), out)
    assert "Its columns are already known" in msg
    assert "header_row=" in msg and '"equals": 1' in msg
    assert "assay column [3, 4]" in msg          # the fixture's two scored columns
    assert "patient col 0" in msg


def test_the_columns_are_derived_only_when_the_gate_is_BLOCKING(tmp_path):
    """Structure costs ~2s per workbook. The passing path must not pay for a message it never emits,
    so a fully-loaded sheet returns '' without ever asking for structure."""
    paper = tmp_path / "paper"
    path = _reactivity_book(paper)
    out = _record(tmp_path, paper)
    ac._append_load_ledger(out, path=str(path), sheet="Suppl Dataset 4b. CD4 T cells",
                           loader="add_table", n_loaded=17, n_available=12)
    calls = []
    real = census.census_tcr_sheets

    def spy(d, with_structure=False):
        calls.append(with_structure)
        return real(d, with_structure=with_structure)

    census.census_tcr_sheets = spy
    try:
        assert ac._reactivity_source_coverage_gap(_paper(out), out) == ""
    finally:
        census.census_tcr_sheets = real
    assert True not in calls, "structure must not be derived on the passing path"


def test_the_inventory_prescribes_the_per_cell_call(tmp_path):
    """Measured 2026-08-25: the finalize gate DOES state these columns, but its message arrives only
    after every add_table call is already made -- one run authored 9 mappings, then saw the columns,
    then overrode. The load decision happens at the INVENTORY, so the column map has to be there,
    exactly like the per-patient sheet-family flag beside it."""
    paper = tmp_path / "paper"
    _reactivity_book(paper)
    txt = ac.survey_sources(str(paper))
    line = next((l for l in txt.splitlines() if "REACTIVITY MATRIX" in l), "")
    assert line, "the inventory must flag a reactivity matrix"
    assert "17 POSITIVE assay call(s)" in line          # the fixture's positive cells
    assert "per_cell" in line and 'assay column [3, 4]' in line
    assert "patient col 0" in line
    assert "equals: 0" in line                          # the C-lite mistake, warned against


def test_the_inventory_is_unchanged_for_a_paper_with_no_reactivity_matrix(tmp_path):
    paper = tmp_path / "paper"
    _track_book(paper)
    assert "REACTIVITY MATRIX" not in ac.survey_sources(str(paper))
