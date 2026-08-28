"""Keskin clone → uniquely resolved target. No guessed mixed-sheet antigen."""
from __future__ import annotations

import re
from pathlib import Path

def _ac():
    from . import agent_core
    return agent_core

_REPORTER_PATTERNS = (
    (re.compile(r"\bH02\b", re.I), "EV_TCR_H02", "immunizing_peptide", "IMP_HNLDLAEKDFMV_20_8eee2c"),
    (re.compile(r"\bF10\b", re.I), "EV_TCR_H02", "immunizing_peptide", "IMP_HNLDLAEKDFMV_20_8eee2c"),
    (re.compile(r"\bEPT12A\b", re.I), "EV_TCR_CD8_EPT12A", "epitope", "EPI_MVNTVAGAMK"),
)
_MULTIMER_CUE = re.compile(r"\btetramer|\bmultimer|\btet-spec\b", re.I)
_ANTIGEN_COL = re.compile(r"\b(antigen|gene|peptide|epitope|neoantigen)\b", re.I)
_CDR3_COL = re.compile(r"cdr3|junction|trav|trbv|clonotype|clone|cell", re.I)
_QUOTED_ANTIGEN = re.compile(
    r"^\S+\s+(.+)-reactive single-cell clonotype\b", re.I
)


def _lane(kind: str) -> str:
    return {"immunizing_peptide": "immunizing_peptides",
            "epitope": "epitopes",
            "candidate": "candidates"}[kind]


def _entity(rec: dict, kind: str, tid: str | None) -> dict | None:
    if not tid:
        return None
    for row in rec.get(_lane(kind)) or []:
        if row.get("paper_local_id") == tid:
            return row
    return None


def _flag_target(flag: dict) -> tuple[str, str] | None:
    kind = flag.get("target_kind")
    tid = (flag.get("immunizing_peptide_paper_id") if kind == "immunizing_peptide"
           else flag.get("epitope_paper_id") if kind == "epitope"
           else flag.get("candidate_paper_id") if kind == "candidate"
           else None)
    if kind and tid:
        return kind, tid
    return None


def _gene_of(rec: dict, kind: str, tid: str) -> str:
    row = _entity(rec, kind, tid)
    return ((row or {}).get("gene_symbol") or "").upper()


def antigen_gene_hits(rec: dict, antigen: str) -> set[str]:
    tokens = {t.upper() for t in _ac()._flag_antigen_tokens(antigen) if len(t) >= 3}
    genes = set()
    for lane in ("immunizing_peptides", "epitopes", "candidates"):
        for row in rec.get(lane) or []:
            g = (row.get("gene_symbol") or "").upper()
            if g and g in tokens:
                genes.add(g)
    return genes


def antigen_is_mixed(rec: dict, antigen: str) -> bool:
    return len(antigen_gene_hits(rec, antigen)) >= 2


def sheet_row_antigen_col(xlsx_path: str | None, sheet: str | None, header_row: int = 1) -> int | None:
    """Return a per-row antigen/gene/peptide column index, or None.

    Mixed sheets stay unbound unless this column exists. Does not treat CDR3/V/J
    columns as antigen.
    """
    if not xlsx_path or not sheet:
        return None
    path = Path(xlsx_path)
    if not path.is_file():
        return None
    try:
        import openpyxl
    except ImportError:
        return None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return None
        ws = wb[sheet]
        rows = ws.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True)
        hdr = next(rows, None)
    finally:
        wb.close()
    if not hdr:
        return None
    hits = []
    for i, cell in enumerate(hdr):
        label = str(cell or "")
        if _ANTIGEN_COL.search(label) and not _CDR3_COL.search(label):
            hits.append(i)
    return hits[0] if len(hits) == 1 else None


def _specificity_for(antigen: str, section_ref: str, sheet: str | None) -> str:
    blob = " ".join(x for x in (antigen, section_ref, sheet) if x)
    if _MULTIMER_CUE.search(blob):
        return "multimer_sort"
    return "not_established"


def _apply_target(clono: dict, kind: str, tid: str, *,
                  specificity: str | None = None, evidence_ref: str | None = None) -> None:
    clono["specific_for_kind"] = kind
    clono["immunizing_peptide_paper_id"] = tid if kind == "immunizing_peptide" else None
    clono["epitope_paper_id"] = tid if kind == "epitope" else None
    clono["candidate_paper_id"] = tid if kind == "candidate" else None
    if evidence_ref:
        clono["validation_evidence_ref"] = evidence_ref
        clono["specificity_evidence"] = "functional_tcr_clone"
    elif specificity:
        # Keep reporter strength if already set.
        if clono.get("specificity_evidence") != "functional_tcr_clone":
            clono["specificity_evidence"] = specificity
    if clono.get("specificity_evidence") != "inferred_association":
        # Bound antigen-sorted rows stay public unless schema-forced.
        if not (len(clono.get("observations") or []) > 1
                and (clono.get("observations_identity") or "not_asserted") == "not_asserted"):
            clono["needs_review"] = False


def _unique_flag_target(rec: dict, *, patient: str, antigen: str,
                        section_ref: str) -> tuple[tuple[str, str] | None, str]:
    tokens = [t.upper() for t in _ac()._flag_antigen_tokens(antigen) if len(t) >= 3]
    if not tokens:
        return None, "no identifying antigen token"
    token_set = set(tokens)
    gene_hits, surface_hits = [], []
    for flag in rec.get("neoantigen_tcr_flags") or []:
        if flag.get("patient_paper_id") != patient:
            continue
        tgt = _flag_target(flag)
        if not tgt:
            continue
        kind, tid = tgt
        gene = _gene_of(rec, kind, tid)
        if gene and gene in token_set:
            gene_hits.append((kind, tid))
        surfaces = " ".join(filter(None, [
            flag.get("quoted_text"), flag.get("section_ref"), gene, tid, section_ref,
        ])).upper()
        if any(t in surfaces for t in token_set):
            surface_hits.append((kind, tid))
    gene_hits = list(dict.fromkeys(gene_hits))
    surface_hits = list(dict.fromkeys(surface_hits))
    if len(gene_hits) == 1:
        return gene_hits[0], ""
    if len(gene_hits) > 1:
        return None, f"{len(gene_hits)} flags match antigen gene tokens — not guessing"
    if len(surface_hits) == 1:
        return surface_hits[0], ""
    if len(surface_hits) > 1:
        return None, f"{len(surface_hits)} flags match antigen tokens — not guessing"
    return None, "no unique flag for this (patient, sheet antigen)"


def bind_clonotype(rec: dict, clono: dict, *, antigen: str, section_ref: str,
                   patient_paper_id: str, sheet: str | None = None,
                   xlsx_path: str | None = None,
                   per_row_antigen: str | None = None) -> str:
    """Mutate one clonotype dict. Returns a skip/bind reason string."""
    ant = per_row_antigen or antigen
    if not ant:
        return "skip: no antigen string"
    if antigen_is_mixed(rec, ant) and not per_row_antigen:
        return "skip: mixed antigen, no per-row column"
    target, why = _ac()._resolve_flag_target(rec, ant)
    if target is None:
        target, why = _unique_flag_target(
            rec, patient=patient_paper_id, antigen=ant, section_ref=section_ref,
        )
    if target is None:
        return f"skip: {why}"
    kind, tid = target
    spec = _specificity_for(ant, section_ref, sheet)
    _apply_target(clono, kind, tid, specificity=spec)
    return f"bound {kind} {tid}"


def reporter_evidence_ids(rec: dict) -> set[str]:
    return {
        e.get("evidence_local_id")
        for e in rec.get("evidence") or []
        if e.get("evidence_local_id") and e.get("assay") == "tcr_reporter"
    }


def apply_reporter_refs(rec: dict, clonotypes: list[dict] | None = None) -> dict:
    """Set validation_evidence_ref only when paper_local_id or quoted_text names H02/F10/EPT12A."""
    rows = clonotypes if clonotypes is not None else (rec.get("tcr_clonotypes") or [])
    known = reporter_evidence_ids(rec)
    report = {"n_linked": 0, "skipped": []}
    for cl in rows:
        blob = f"{cl.get('paper_local_id') or ''} {cl.get('quoted_text') or ''}"
        hits = []
        for pat, ev_id, kind, tid in _REPORTER_PATTERNS:
            if pat.search(blob):
                hits.append((ev_id, kind, tid))
        hits = list(dict.fromkeys(hits))
        if not hits:
            continue
        if len(hits) != 1:
            report["skipped"].append({
                "paper_local_id": cl.get("paper_local_id"),
                "reason": f"names {len(hits)} reporter ids — not guessing",
            })
            continue
        ev_id, kind, tid = hits[0]
        if ev_id not in known:
            report["skipped"].append({
                "paper_local_id": cl.get("paper_local_id"),
                "reason": f"{ev_id} not a tcr_reporter evidence_local_id in the record",
            })
            continue
        _apply_target(cl, kind, tid, evidence_ref=ev_id)
        report["n_linked"] += 1
    if report["n_linked"] == 0 and not report["skipped"]:
        report["skipped"].append({
            "paper_local_id": None,
            "reason": "no clonotype paper_local_id or quoted_text names H02, F10, or EPT12A",
        })
    return report


def antigen_from_clonotype(cl: dict) -> str:
    qt = cl.get("quoted_text") or ""
    m = _QUOTED_ANTIGEN.match(qt)
    if m:
        return m.group(1).strip()
    return ""


def bind_record(rec: dict, *, xlsx_by_section: dict[str, tuple[str, str]] | None = None) -> dict:
    """Bind every clonotype in `rec`. xlsx_by_section maps section_ref -> (path, sheet)."""
    xlsx_by_section = xlsx_by_section or {}
    report = {
        "n_bound": 0, "n_skipped": 0, "n_mixed": 0, "n_reporter": 0,
        "skipped": [], "mixed_sheets": [],
    }
    mixed_seen: set[str] = set()
    for cl in rec.get("tcr_clonotypes") or []:
        antigen = antigen_from_clonotype(cl)
        section_ref = cl.get("section_ref") or ""
        patient = cl.get("patient_paper_id") or ""
        xinfo = xlsx_by_section.get(section_ref)
        path, sheet = xinfo if xinfo else (None, None)
        per_row = None
        mixed = antigen_is_mixed(rec, antigen)
        if mixed:
            col = sheet_row_antigen_col(path, sheet)
            if col is None:
                report["n_mixed"] += 1
                report["n_skipped"] += 1
                mixed_seen.add(section_ref or antigen)
                continue
            # Per-row bind is only possible at load time (cell grouping). Offline,
            # mixed sheets without a stored per-clone antigen stay skipped.
            report["n_mixed"] += 1
            report["n_skipped"] += 1
            mixed_seen.add(section_ref or antigen)
            continue
        why = bind_clonotype(
            rec, cl, antigen=antigen, section_ref=section_ref,
            patient_paper_id=patient, sheet=sheet, xlsx_path=path,
            per_row_antigen=per_row,
        )
        if why.startswith("bound"):
            report["n_bound"] += 1
        else:
            report["n_skipped"] += 1
            report["skipped"].append({"paper_local_id": cl.get("paper_local_id"),
                                      "section_ref": section_ref, "reason": why})
    report["mixed_sheets"] = sorted(mixed_seen)
    rep = apply_reporter_refs(rec)
    report["n_reporter"] = rep["n_linked"]
    report["reporter_skipped"] = rep["skipped"]
    return report
